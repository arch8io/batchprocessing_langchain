import os
from dotenv import load_dotenv
from langchain_openai import ChatOpenAI
from openpyxl import load_workbook

load_dotenv()

# Point to OpenRouter using a free model
llm = ChatOpenAI(
    # Option 1: Meta Llama 3.3 70B Free
    # model="meta-llama/llama-3.3-70b-instruct:free",
    
    # Option 2: Google Gemini Flash 2.0 Free (Alternative)
    # model="google/gemini-2.0-flash-exp:free",
    
    model="nvidia/nemotron-3-ultra-550b-a55b:free",

    openai_api_key=os.getenv("OPENROUTER_API_KEY"),
    openai_api_base="https://openrouter.ai/api/v1",
    temperature=0,
)

# Excel file path
file_path_feedback = 'data/Customer_Sentiment.xlsx'
wb = load_workbook(file_path_feedback)
sheet = wb.active

for row in sheet.iter_rows(values_only=True, min_row=2):
    feedback = row[8]
    result = llm.invoke(f"""You need to classify the reason for the user's feedback.
The categories include: Price too high, Insufficient after-sales support, Poor product experience, Other.
Response format: Classification result: xx.
The user's issue is: {feedback}""")
    print(feedback,result.content)
